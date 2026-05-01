#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import subprocess
from pathlib import Path

from docx import Document  # type: ignore
from openpyxl import load_workbook  # type: ignore

try:
    from PyPDF2 import PdfReader  # type: ignore
except Exception:  # pragma: no cover
    PdfReader = None


VISUAL_HINTS = [
    "logo",
    "svg",
    "矢量",
    "草图",
    "图标",
    "标志",
    "视觉",
]
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".svg"}


def read_prompt(prompt_path: Path) -> str:
    text = prompt_path.read_text(encoding="utf-8")
    parts = text.split("\n\n", 1)
    return parts[1].strip() if len(parts) > 1 else text.strip()


def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def read_xlsx(path: Path) -> str:
    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append(" | ".join(values))
        lines.append("")
    return "\n".join(lines).strip()


def read_csv_like(path: Path) -> str:
    lines: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as fh:
        reader = csv.reader(fh)
        for row in reader:
            lines.append(" | ".join(row))
    return "\n".join(lines)


def read_pdf(path: Path) -> str:
    if PdfReader is None:
        return f"[PDF parser unavailable for {path.name}]"
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def read_source(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return read_text_file(path)
    if suffix == ".json":
        return json.dumps(json.loads(read_text_file(path)), ensure_ascii=False, indent=2)
    if suffix in {".csv", ".tsv"}:
        return read_csv_like(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    if suffix == ".pdf":
        return read_pdf(path)
    return f"[Unsupported file type: {path.name}]"


def build_context(input_dir: Path, max_chars: int) -> tuple[str, list[dict[str, str]]]:
    sections: list[str] = []
    files_meta: list[dict[str, str]] = []
    used = 0
    for path in sorted(p for p in input_dir.rglob("*") if p.is_file()):
        content = read_source(path).strip()
        files_meta.append({"name": path.name, "path": str(path)})
        if not content:
            continue
        block = f"## Source File: {path.name}\n\n{content}\n"
        if used + len(block) > max_chars:
            remaining = max(0, max_chars - used)
            if remaining > 0:
                sections.append(block[:remaining])
            break
        sections.append(block)
        used += len(block)
    return "\n".join(sections).strip(), files_meta


def is_visual_prompt(prompt: str, files_meta: list[dict[str, str]]) -> bool:
    lowered = prompt.lower()
    has_design_intent = any(token in lowered for token in ["设计", "绘制", "画", "重做", "优化", "改一下", "再设计", "说明设计思路"])
    has_visual_hint = any(token in lowered for token in VISUAL_HINTS)
    has_image_source = any(Path(item.get("name", "")).suffix.lower() in IMAGE_SUFFIXES for item in files_meta)
    return (has_design_intent and has_visual_hint) or (has_design_intent and has_image_source)


def build_codex_prompt(prompt: str, source_context: str, manifest_path: Path, *, output_mode: str) -> str:
    if output_mode == "svg":
        return "\n".join(
            [
                "You are executing a reviewed workflow job.",
                "Use the provided source files as the primary evidence base.",
                "Generate one clean standalone SVG logo concept.",
                "The final answer must be raw SVG markup only.",
                "Do not wrap the SVG in Markdown fences.",
                "Do not add any explanation before or after the SVG.",
                "Keep the SVG vector-only and self-contained, with no external assets.",
                "Prefer a professional brand-mark style suitable for a strategy draft.",
                "",
                f"Manifest path: {manifest_path}",
                "",
                "# User Prompt",
                prompt.strip(),
                "",
                "# Source Materials",
                source_context.strip() or "[No readable source files were found.]",
            ]
        ).strip() + "\n"

    return "\n".join(
        [
            "You are executing a reviewed workflow job.",
            "Use the provided source files as the primary evidence base.",
            "Return a polished final answer in Markdown.",
            "When the prompt asks for a report, make it structured and detailed.",
            "When the prompt asks for naming or strategy work, provide concrete options, rationale, and implementation steps.",
            "Do not mention missing files unless the source section is empty.",
            "",
            f"Manifest path: {manifest_path}",
            "",
            "# User Prompt",
            prompt.strip(),
            "",
            "# Source Materials",
            source_context.strip() or "[No readable source files were found.]",
        ]
    ).strip() + "\n"


def extract_svg_markup(text: str) -> str:
    stripped = text.strip()
    if stripped.startswith("```"):
        lines = stripped.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        stripped = "\n".join(lines).strip()
    start = stripped.find("<svg")
    end = stripped.rfind("</svg>")
    if start == -1 or end == -1:
        raise RuntimeError("codex did not return SVG markup")
    end += len("</svg>")
    return stripped[start:end] + "\n"


def run_codex(
    codex_bin: str,
    model: str,
    workdir: Path,
    prompt_text: str,
    output_path: Path,
    trace_path: Path,
) -> str:
    cmd = [
        codex_bin,
        "exec",
        "--skip-git-repo-check",
        "--sandbox",
        "read-only",
        "--output-last-message",
        str(output_path),
        "--cd",
        str(workdir),
        "-",
    ]
    if model.strip():
        cmd.extend(["--model", model.strip()])

    proc = subprocess.run(
        cmd,
        input=prompt_text,
        text=True,
        capture_output=True,
        check=False,
    )
    trace = {
        "command": cmd,
        "returncode": proc.returncode,
        "stdout": proc.stdout,
        "stderr": proc.stderr,
    }
    trace_path.write_text(json.dumps(trace, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or proc.stdout.strip() or "codex exec failed")
    if not output_path.exists():
        raise RuntimeError("codex exec finished without writing the final message output")
    return model


def main() -> int:
    parser = argparse.ArgumentParser(description="Execute prompt jobs with local Codex and attached source files")
    parser.add_argument("--prompt", required=True)
    parser.add_argument("--inputs", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--runner", default=os.getenv("PROMPT_EXECUTOR_RUNNER", "codex"))
    parser.add_argument("--model", default=os.getenv("PROMPT_EXECUTOR_MODEL", "gpt-5.4"))
    parser.add_argument("--visual-model", default=os.getenv("PROMPT_EXECUTOR_VISUAL_MODEL", "gpt-5.5"))
    parser.add_argument("--codex-bin", default=os.getenv("CODEX_BIN", "codex"))
    parser.add_argument("--max-source-chars", type=int, default=30000)
    args = parser.parse_args()

    prompt_path = Path(args.prompt)
    input_dir = Path(args.inputs)
    output_dir = Path(args.output)
    manifest_path = Path(args.manifest)
    output_dir.mkdir(parents=True, exist_ok=True)

    prompt = read_prompt(prompt_path)
    source_context, files_meta = build_context(input_dir, args.max_source_chars)
    output_mode = "svg" if is_visual_prompt(prompt, files_meta) else "markdown"
    selected_model = args.visual_model if output_mode == "svg" else args.model
    prompt_text = build_codex_prompt(prompt, source_context, manifest_path, output_mode=output_mode)

    rendered_prompt_path = output_dir / "executor_prompt.txt"
    rendered_prompt_path.write_text(prompt_text, encoding="utf-8")
    trace_path = output_dir / "executor_trace.json"
    result_path = output_dir / ("result.svg" if output_mode == "svg" else "result.md")

    runner = args.runner.lower().strip()
    if runner != "codex":
        raise RuntimeError(f"unsupported runner: {runner}")

    actual_model = selected_model
    model_fallback = ""
    try:
        actual_model = run_codex(
            codex_bin=args.codex_bin,
            model=selected_model,
            workdir=manifest_path.parent,
            prompt_text=prompt_text,
            output_path=result_path,
            trace_path=trace_path,
        )
    except RuntimeError as exc:
        detail = str(exc)
        if (
            output_mode == "svg"
            and args.visual_model.strip()
            and args.visual_model.strip() != args.model.strip()
            and "requires a newer version of Codex" in detail
        ):
            model_fallback = args.model
            actual_model = run_codex(
                codex_bin=args.codex_bin,
                model=args.model,
                workdir=manifest_path.parent,
                prompt_text=prompt_text,
                output_path=result_path,
                trace_path=trace_path,
            )
        else:
            raise

    if output_mode == "svg":
        svg = extract_svg_markup(result_path.read_text(encoding="utf-8", errors="ignore"))
        result_path.write_text(svg, encoding="utf-8")
        summary_path = output_dir / "result.md"
        summary_path.write_text(
            "\n".join(
                [
                    "# SVG Logo Draft",
                    "",
                    "Generated a vector logo draft based on the reviewed prompt and attached source files.",
                    "",
                    "- Primary artifact: `result.svg`",
                    "",
                ]
            )
            + "\n",
            encoding="utf-8",
        )

    metadata = {
        "runner": runner,
        "model": selected_model,
        "actual_model": actual_model,
        "default_model": args.model,
        "visual_model": args.visual_model,
        "model_fallback": model_fallback,
        "codex_bin": args.codex_bin,
        "output_mode": output_mode,
        "prompt_path": str(prompt_path),
        "manifest_path": str(manifest_path),
        "source_files": files_meta,
        "result_path": str(result_path),
        "executor_prompt_path": str(rendered_prompt_path),
        "executor_trace_path": str(trace_path),
    }
    (output_dir / "result.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
